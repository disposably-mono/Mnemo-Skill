"""PPTX ingestion via python-pptx: one Chunk per slide, with table/chart/notes."""

from __future__ import annotations

import logging
from pathlib import Path

from mnemo.ingest import Chunk

_logger = logging.getLogger(__name__)


def _format_chart_value(value) -> str:
    """Render chart values without unnecessary floating-point noise."""
    if value is None:
        return ""
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return format(value, ".12g")
    return str(value).strip()


def _extract_pptx_chart(chart) -> str:
    """Render accessible chart categories and series as a relational table."""
    title = ""
    try:
        if chart.has_title:
            title = chart.chart_title.text_frame.text.strip()
    except Exception:
        _logger.debug("reading pptx chart title failed", exc_info=True)

    series = list(chart.series)
    names = [
        str(item.name).strip() or f"Series {index}"
        for index, item in enumerate(series, start=1)
    ]
    try:
        categories = [
            str(getattr(category, "label", category)).strip()
            for category in chart.plots[0].categories
        ]
    except Exception:
        _logger.debug("reading pptx chart categories failed", exc_info=True)
        categories = []

    lines = ["Chart:"]
    if title:
        lines.append(f"Title: {title}")
    if categories and series:
        lines.append("Category | " + " | ".join(names))
        values = [list(item.values) for item in series]
        for index, category in enumerate(categories):
            row = [
                _format_chart_value(series_values[index])
                if index < len(series_values)
                else ""
                for series_values in values
            ]
            lines.append(f"{category} | " + " | ".join(row))
    elif names:
        lines.append("Series: " + " | ".join(names))
    return "\n".join(lines) if len(lines) > 1 else ""


def _shape_parts(shapes) -> list[str]:
    """Extract text and structured content recursively, including group shapes."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    parts: list[str] = []
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            parts.extend(_shape_parts(shape.shapes))
            continue
        if shape.has_text_frame and shape.text_frame.text.strip():
            parts.append(shape.text_frame.text)
        elif shape.has_table:
            rows = [
                " | ".join(cell.text.strip() for cell in row.cells)
                for row in shape.table.rows
            ]
            table = "\n".join(row for row in rows if row.strip(" |"))
            if table:
                parts.append(table)
        elif shape.has_chart:
            chart = _extract_pptx_chart(shape.chart)
            if chart:
                parts.append(chart)
    return parts


def ingest_pptx(path: Path) -> list[Chunk]:
    """Ingest a PPTX: one Chunk per non-empty slide, with speaker notes appended."""
    from pptx import Presentation

    chunks: list[Chunk] = []
    prs = Presentation(str(path))
    for index, slide in enumerate(prs.slides, start=1):
        parts: list[str] = []
        parts.extend(_shape_parts(slide.shapes))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"Speaker notes:\n{notes}")
        text = "\n".join(parts).strip()
        if text:
            chunks.append(Chunk(text=text, source=f"{path.name} slide {index}"))
    return chunks
