"""Tests for DOCX ingestion (mnemo/ingest/docx.py)."""

from mnemo.ingest import ingest


def test_docx_is_read_as_a_single_chunk(tmp_path):
    import docx

    path = tmp_path / "notes.docx"
    document = docx.Document()
    document.add_paragraph("The mitochondria is the powerhouse of the cell.")
    document.save(str(path))

    chunks = ingest(path)

    assert len(chunks) == 1
    assert "powerhouse" in chunks[0].text
    assert chunks[0].source == "notes.docx"


def test_empty_docx_yields_no_chunks(tmp_path):
    import docx

    path = tmp_path / "empty.docx"
    docx.Document().save(str(path))

    assert ingest(path) == []
