"""Tests for PPTX ingestion (mnemo/ingest/pptx.py)."""

from mnemo.ingest import ingest


def test_pptx_yields_one_chunk_per_slide_with_shape_text(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    path = tmp_path / "deck.pptx"
    prs = Presentation()
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    box.text_frame.text = "Newton's first law: inertia."
    prs.save(str(path))

    chunks = ingest(path)

    assert len(chunks) == 1
    assert "inertia" in chunks[0].text
    assert chunks[0].source == "deck.pptx slide 1"


def test_pptx_includes_table_and_speaker_notes(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    path = tmp_path / "deck.pptx"
    prs = Presentation()
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    graphic_frame = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(1))
    table = graphic_frame.table
    table.cell(0, 0).text = "Planet"
    table.cell(0, 1).text = "Radius"
    table.cell(1, 0).text = "Earth"
    table.cell(1, 1).text = "6371 km"
    slide.notes_slide.notes_text_frame.text = "Remember the order of the planets."
    prs.save(str(path))

    chunks = ingest(path)

    assert len(chunks) == 1
    assert "Planet | Radius" in chunks[0].text
    assert "Speaker notes:" in chunks[0].text
    assert "Remember the order" in chunks[0].text


def test_pptx_chart_is_rendered_as_category_table(tmp_path):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    path = tmp_path / "deck.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    chart_data = CategoryChartData()
    chart_data.categories = ["Town A", "Town B"]
    chart_data.add_series("2020", (10, 20))
    slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(5), Inches(4), chart_data
    )
    prs.save(str(path))

    chunks = ingest(path)

    assert len(chunks) == 1
    assert "Chart:" in chunks[0].text
    assert "Category | 2020" in chunks[0].text
    assert "Town A | 10" in chunks[0].text
    assert "Town B | 20" in chunks[0].text


def test_pptx_empty_slide_is_skipped(tmp_path):
    from pptx import Presentation

    path = tmp_path / "deck.pptx"
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    prs.save(str(path))

    assert ingest(path) == []
